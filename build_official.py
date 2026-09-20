# -*- coding: utf-8 -*-
"""把我们的内容按北邮《立项申请书》官方模板重排"""
import shutil, datetime, os
from docx import Document
from docx.shared import Pt, Emu
from docx.oxml.ns import qn

TPL = "通用大模型与场景化 Agent 协同开发与应用研究-宋科翰.docx"
OUT = "基于大模型的模拟电子技术解题诊断与变式练习系统-宋林睿.docx"

NAME = "基于大模型的模拟电子技术解题诊断与变式练习系统"
ENAME = ("A Large Language Model-Based System for Solution Diagnosis "
         "and Variational Practice in Analog Electronics")

def _prepare():
    """目标文件被 Word 占用时，自动改用带时间戳的副本"""
    try:
        shutil.copyfile(TPL, OUT)
        return OUT
    except PermissionError:
        import time
        alt = OUT.replace(".docx", "_" + time.strftime("%m%d_%H%M") + ".docx")
        shutil.copyfile(TPL, alt)
        print("LOCKED:", OUT, "->", alt)
        return alt


OUT = _prepare()
d = Document(OUT)


# ---------- 工具函数 ----------
def set_runs(p, text, size=None, bold=None):
    """保留段落格式，只替换文字"""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return run


def set_para(p, text):
    """只改文字，完全不动原有格式"""
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        run = p.add_run(str(text))
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        run = p.add_run(str(text))
        run.font.size = Pt(10.5)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


paras = d.paragraphs
tables = d.tables

# ================= 封面 =================
for p in paras:
    t = p.text.strip()
    if t.startswith("项目来源") or t.startswith("揭榜挂帅类（企业）") or t.startswith("国际专项") \
       or t.startswith("科创融合"):
        set_para(p, t.replace("☑", "□"))
    if "自主探索类" in t:
        set_para(p, t.replace("自主探索类□", "自主探索类☑"))
    elif t.startswith("项目名称:") or t.startswith("项目名称："):
        set_para(p, f"项目名称:  {NAME}")
    elif t.startswith("项目名称（英文）"):
        set_para(p, f"项目名称（英文）： {ENAME}")
    elif t.startswith("项目依托学院"):
        set_para(p, "项目依托学院：     人工智能学院                               ")
    elif t.startswith("项目负责人"):
        set_para(p, "项目负责人：       宋林睿                                  ")
    elif t.startswith("联系电话"):
        set_para(p, "联系电话：         15201175336                                ")
    elif t.startswith("E-mail") and "skh" in t:
        set_para(p, "E-mail：           snlxrain@qq.com                            ")
    elif t.startswith("指导教师"):
        set_para(p, "指导教师：         望育梅、迟国轩                            ")
    elif t.startswith("E-mail") and "ymwang" in t:
        set_para(p, "E-mail：  ymwang@bupt.edu.cn / chiguoxuan@bupt.edu.cn        ")
    elif t.startswith("填报时间"):
        now = datetime.datetime.now()
        set_para(p, f"填报时间:{now.year}年  {now.month}  月  {now.day} 日")

# ================= 一、基本情况表 =================
T = tables[0]
R = T.rows

# R0 项目名称
set_cell(R[0].cells[1], NAME)

# R1 负责人：学号 / 学院 / 手机号
set_cell(R[1].cells[1], "宋林睿")
set_cell(R[1].cells[4], "2025212007")
set_cell(R[1].cells[6], "人工智能学院")
set_cell(R[1].cells[8], "15201175336")

# R2 负责人：专业 / 班级 / 邮箱
set_cell(R[2].cells[1], "宋林睿")
set_cell(R[2].cells[4], "信息工程")
set_cell(R[2].cells[6], "2025219104")
set_cell(R[2].cells[8], "snlxrain@qq.com")

# R3/R4 指导教师一：望育梅
set_cell(R[3].cells[1], "望育梅")
set_cell(R[3].cells[4], "副教授")
set_cell(R[3].cells[6], "人工智能学院")
set_cell(R[3].cells[8], "18515155866")
set_cell(R[4].cells[1], "望育梅")
set_cell(R[4].cells[4], "副教授")
set_cell(R[4].cells[6], "人工智能学院")
set_cell(R[4].cells[8], "ymwang@bupt.edu.cn")

# R5/R6 指导教师二：迟国轩
set_cell(R[5].cells[1], "迟国轩")
set_cell(R[5].cells[4], "助理教授")
set_cell(R[5].cells[6], "人工智能学院")
set_cell(R[5].cells[8], "18811521756")
set_cell(R[6].cells[1], "迟国轩")
set_cell(R[6].cells[4], "助理教授")
set_cell(R[6].cells[6], "人工智能学院")
set_cell(R[6].cells[8], "chiguoxuan@bupt.edu.cn")

# R7 检索关键词
set_cell(R[7].cells[1], "大语言模型；模拟电子技术；解题过程诊断；变式练习；智能辅导")

# R8 表头保留；R9 成员薛群辉；R10-R12 清空（第0列是纵向合并标签，不能动）
set_cell(R[9].cells[1], "薛群辉")
set_cell(R[9].cells[2], "人工智能学院")
set_cell(R[9].cells[4], "信息工程")
set_cell(R[9].cells[5], "2025219105")
set_cell(R[9].cells[6], "2025211919")
set_cell(R[9].cells[7], "")
set_cell(R[9].cells[8], "2370839713@qq.com")
for i in (10, 11, 12):
    for j in range(1, 9):
        set_cell(R[i].cells[j], "")

# R13 团队主要成员介绍
set_cell(R[13].cells[1],
         "宋林睿：人工智能学院信息工程专业2025级本科生，已修电子技术基础（模拟电子技术）、高等数学、"
         "程序设计基础等课程，掌握Python编程与大模型接口调用，负责项目总体设计、考点图谱与"
         "解题诊断逻辑开发。\n"
         "薛群辉：人工智能学院信息工程专业2025级本科生，具备Python编程与数据处理基础，了解机器学习基本"
         "流程，负责题库建设与数据标注、前端界面开发与试用测试。")

# R14 指导教师承担科研课题情况
set_cell(R[14].cells[1],
         "望育梅老师主持或参与多项国家自然科学基金项目、国家重点研发计划、北京市自然科学基金等"
         "国家和省部级项目，正研科研项目有：\n"
         "在轨微云数字样机研发；面向卫星互联网星座的云网融合关键技术研究；"
         "在轨微云集成测试验证评估平台研发等。\n"
         "迟国轩老师主持或参与国家自然科学基金、国家重点研发计划子课题、中国博士后科学基金等"
         "国家级及省部级项目，正研科研项目有：\n"
         "主持项目：国家自然科学基金青年科学基金项目（C类），项目编号62402276，资助额度30万元；"
         "中国博士后科学基金面上资助，项目编号2025M771498，资助额度8万元。\n"
         "参与项目：国家自然科学基金重点项目、面上项目；国家重点研发计划子课题。\n"
         "研究方向关联项目：面向无线系统的生成式人工智能（如RF-Diffusion射频信号生成）；"
         "射频与无线感知（Wi-Fi/毫米波感知）；具身导航与空间智能；通信感知一体化。")

# R15 指导教师对本项目支持情况
set_cell(R[15].cells[1],
         "1.定期召开组会并进行监督指导，通过线上线下方式与组内成员沟通交流；\n"
         "2.为项目提供必要的技术指导，在考点体系、常见错因判据与题目来源方面给予支持，"
         "指导项目的开发和落地；\n"
         "3.进行研究方法、理论学习、技术实现等方面的引导，并对系统的诊断结果进行审核把关；\n"
         "4.在可能的情况下，提供必要的实验条件和经费支持。")

# ================= 二、立项依据 =================
CONTENT = [
 ("sub", "（一）创意缘起与价值定位"),
 ("body", "模拟电子技术是工科电类专业的核心基础课，也是公认难学、挂科率较高的课程。课程概念抽象、"
          "分析方法繁多，学生的主要学习方式是课后做题，但普遍存在“答案看得懂、自己做不对”的现象——"
          "能看懂标准答案，轮到自己动手就卡在中间某一步。这说明学生的困难不在最终结论，而在解题过程的中间环节。"),
 ("body", "现有题库与练习类软件大多只做最终答案比对，既不指出错在哪一步，也不区分错误类型，"
          "学生只能反复做同类错题，效率很低。而模电的错因具有明显的类型化与规律性，例如小信号等效电路画错、"
          "静态工作点计算错误、近似条件误用、反馈极性判断错误等，具备建模与自动识别的可行性。"),
 ("body", "本项目拟构建“练习—诊断—变式再练”的智能学习系统：学生分步提交自己的解题过程，"
          "系统定位首次出错的步骤并判定错因，给出针对该错因的讲解，再生成同考点、同难度的变式练习题，形成学习闭环。"),
 ("body", "本项目的研究价值主要体现在三个方面："),
 ("body", "第一，把学习诊断的粒度从“对/错”细化到“哪一步错、为什么错”，使反馈真正落在学生的困难处，"
          "而不是让学生在一堆同类题里盲目重复。"),
 ("body", "第二，把模电的错因经验数字化、模型化，沉淀为可复用的考点图谱与错因标注数据集，"
          "为课程教学改革与后续教学研究提供数据基础。"),
 ("body", "第三，用“大模型生成 + 符号计算校验”的双通道方案保证解答与讲解的可靠性，"
          "为同类工科基础课程的智能辅导提供可迁移的技术路径。"),

 ("sub", "（二）研究核心与实施要点"),
 ("subsub", "主要工作内容"),
 ("body", "（1）构建模电考点与知识点图谱：覆盖二极管与整流电路、BJT/FET放大电路、差分放大电路、"
          "频率响应、负反馈、集成运放、功率放大、直流电源等模块，梳理知识点与考点的依赖关系，目标不少于60个知识点。"),
 ("body", "（2）建设结构化题库：自编或改编题目不少于300道，每题标注知识点、难度、标准解题步骤与关键判据，并附参考答案。"),
 ("body", "（3）实现解题步骤诊断：支持学生分步输入或公式识别，将学生解题路径与标准路径逐步对齐，"
          "定位首次出错的步骤并归类错因（目标不少于8类）。"),
 ("body", "（4）实现变式题生成：围绕同一知识点与难度，通过改参数、改问法、改电路拓扑自动生成变式题，"
          "并由指导教师抽检合格率。"),
 ("body", "（5）构建掌握度画像与推荐：依据诊断结果统计各知识点的掌握情况，推送薄弱环节的练习与讲解。"),
 ("body", "（6）完成系统集成与试用：开发Web端界面，在课程学习小组或习题课中开展小规模试用与效果评测。"),
 ("subsub", "现实难题攻克"),
 ("body", "重点解决三个问题：一是大模型解模电题时可能算错或讲解不可靠，通过引入SymPy符号计算对电路方程"
          "求解结果进行自动校验，对关键题目辅以LTspice/Multisim仿真交叉验证，形成“模型生成+数值校验+教师抽检”"
          "的三重机制；二是学生手写电路图自动识别难度大，一期将作答方式收敛为分步文字与公式输入、选择题式作答，"
          "保证核心功能按期交付，电路图识别作为后续拓展目标；三是题库与标注工作量大，先聚焦BJT放大电路、"
          "集成运放等核心章节形成完整流程，再横向扩展。"),

 ("sub", "（三）项目创新亮点与独特优势"),
 ("body", "任务粒度创新：把学习诊断从“判断答案对错”推进到“定位错在第几步、属于哪一类错因”，"
          "直击“看答案能懂、自己做不对”的真实痛点。"),
 ("body", "方法创新：大模型生成与符号计算校验双通道结合，用可验证的求解结果约束生成内容，"
          "显著提升解答与讲解的可靠性，避免“算错还讲得头头是道”。"),
 ("body", "闭环设计：诊断结果直接驱动变式题生成与薄弱点推荐，形成“诊断—讲解—再练”的完整学习闭环。"),
 ("body", "数据价值：沉淀模电考点图谱与错因标注数据集，具备复用价值，可支撑后续教学研究与同类课程迁移。"),
 ("body", "低成本可复制：全流程基于开源模型与开源工具链，普通学生设备即可运行，便于在校内推广。"),

 ("sub", "（四）系统架构与实施路径"),
 ("subsub", "开发环境与模型配置"),
 ("body", "开发环境以团队现有计算机为基础，使用Python实现数据处理、诊断逻辑与服务接口；"
          "大模型通过统一适配层接入，记录模型标识、版本、参数与调用消耗；不将从零训练通用大模型作为本项目任务。"),
 ("body", "题目与知识资源以自编题目、改编题目和指导教师自有题库为主，标注来源，不整册扫描或复制教材习题；"
          "涉及学生作答数据的采集遵循匿名化原则，仅用于本项目的研究与测试。"),
 ("body", "工具环境包括公式识别、符号计算（SymPy）、仿真校验（LTspice/Multisim）与数据统计等接口，"
          "统一参数与返回格式，并配置超时、错误返回与日志记录。"),
 ("subsub", "软件架构"),
 ("body", "软件系统按交互层、诊断层、知识与工具层、评价层组织，各模块通过统一接口连接，具体如下："),
 ("body", "交互与任务管理模块：接收学生分步作答内容（文字、公式或选择题式作答），为每次练习生成唯一编号，"
          "展示诊断进度与历史错题记录，并对缺少关键条件的作答请求补充说明。"),
 ("body", "诊断与生成模块：将学生解题路径与标准步骤逐步对齐，定位首次偏离的步骤并判定错因类型；"
          "随后调用变式题生成模块，输出同考点、同难度的练习题与针对性讲解。"),
 ("body", "知识与工具模块：管理知识点图谱、结构化题库与错因判据；符号计算模块对电路方程求解结果做自动校验，"
          "仿真模块对关键题目做交叉验证，校验不通过的题目不进入正式题库。"),
 ("body", "评价与服务模块：以Web页面呈现诊断报告、错因统计与薄弱点推荐；建立“仅答案判定”与"
          "“步骤诊断+变式练习”两组对照，记录诊断准确率、变式题合格率与学生练习正确率的变化。"),

 ("sub", "（五）实施进度规划"),
 ("body", "第一阶段（2026年9月—10月）：完成教学现状与需求调研，设计考点与知识点图谱，"
          "确定题目收集与改编方案，形成需求文档及评价方案。"),
 ("body", "第二阶段（2026年11月—12月）：完成题库结构化建设，标注标准解题步骤与错因判据，"
          "优先完成BJT放大电路、集成运放等核心章节。"),
 ("body", "第三阶段（2027年1月—2月）：利用寒假集中攻关，完成大模型解题与SymPy符号校验模块开发，"
          "验证答案与讲解的正确性。"),
 ("body", "第四阶段（2027年3月—4月）：完成解题步骤诊断模型与变式题生成模块开发，完成系统集成与内部测试，"
          "形成可演示原型及首轮诊断结果。"),
 ("body", "第五阶段（2027年5月）：在课程学习小组或习题课中开展小规模试用，完成指标评测与系统迭代。"),
 ("body", "第六阶段（2027年6月）：撰写结题报告，申请软件著作权，制作演示视频并整理结题材料。"),

 ("sub", "（六）现有条件"),
 ("body", "团队成员为人工智能学院2025级本科生，已修或在修电子技术基础（模拟电子技术）、高等数学、"
          "程序设计基础等课程，熟悉放大电路、频率响应、反馈与集成运放等核心内容，理解课程的主要考点与常见错误。"),
 ("body", "成员具备Python编程与数据处理基础，能够调用大模型接口、使用符号计算库并完成基础Web开发，"
          "可承担资料整理、原型开发、实验设计与成果展示等工作。"),
 ("body", "指导教师长期从事相关课程教学与科研工作，可提供考点体系、题目来源与错因判据，"
          "并对诊断结果进行审核把关，在研究方法、系统设计与技术实现方面提供指导。"),
 ("body", "项目以软件研发为主，优先利用现有学习与开发条件开展原型开发和小规模验证，"
          "模型调用与必要计算服务纳入经费统筹；尚未取得的题目资源与算力资源不作为既有成果，按研究进度逐步落实。"),

 ("sub", "（七）预期输出与考核指标"),
 ("body", "诊断与练习系统原型：完成1套面向模拟电子技术课程学习的Web原型，支持分步作答、错因诊断、"
          "变式练习生成、错题本与薄弱点推荐，保留诊断依据与执行记录，提供可复现的典型案例。"),
 ("body", "题库与数据集：建设结构化题库不少于300道、覆盖知识点不少于60个；形成模电错因标注数据集1份。"),
 ("body", "指标目标：错因诊断（首次出错步骤定位与错因分类）准确率不低于80%；生成的变式题经指导教师"
          "抽检合格率不低于80%；参与试用学生不少于50人，试用后认为“有助于找到自己出错原因”的比例不低于80%。"
          "最终如实报告测试结果及失败案例。"),
 ("body", "研究与开发文档：形成结题报告1份、试用与评测报告1份，以及系统说明与使用文档；申请软件著作权1项。"),
]

# 定位“二、立项依据”与“三、经费概算”之间的段落
i_start = i_end = None
for i, p in enumerate(paras):
    if p.text.strip() == "二、立项依据":
        i_start = i
    if p.text.strip().startswith("三、经费概算") and i_start is not None:
        i_end = i
        break
pool = paras[i_start + 1:i_end]

for idx, (kind, text) in enumerate(CONTENT):
    if idx < len(pool):
        p = pool[idx]
    else:  # 段落池不够时，克隆上一段并追加到池中
        import copy
        from docx.text.paragraph import Paragraph
        new_el = copy.deepcopy(pool[-1]._element)
        pool[-1]._element.addnext(new_el)
        p = Paragraph(new_el, pool[-1]._parent)
        pool.append(p)
    set_runs(p, text, size=12, bold=(kind in ("sub", "subsub")))
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(2)
    pf.space_before = Pt(6) if kind in ("sub", "subsub") else Pt(0)
    if kind == "body":
        pf.first_line_indent = Emu(266700)
    else:
        pf.first_line_indent = None

for p in pool[len(CONTENT):]:
    p._element.getparent().remove(p._element)

# ================= 三、经费概算 =================
set_cell(tables[1].rows[0].cells[0], "（一）项目总经费： 2000元")

# ================= 四、成员分工 =================
T3 = tables[2]
set_cell(T3.rows[1].cells[0], "宋林睿")
set_cell(T3.rows[1].cells[1], "已修电子技术基础（模拟电子技术）、高等数学、程序设计基础等课程，"
                              "掌握Python编程与大模型接口调用")
set_cell(T3.rows[1].cells[2], "负责项目总体设计、考点图谱与解题诊断逻辑开发、进度协调与研究文档统筹")
set_cell(T3.rows[1].cells[3], "200天")
set_cell(T3.rows[1].cells[4], "宋林睿")
set_cell(T3.rows[2].cells[0], "薛群辉")
set_cell(T3.rows[2].cells[1], "具备Python编程与数据处理基础，了解机器学习基本流程")
set_cell(T3.rows[2].cells[2], "负责题库建设与数据标注、前端界面开发、试用测试与结果分析")
set_cell(T3.rows[2].cells[3], "200天")
set_cell(T3.rows[2].cells[4], "薛群辉")
for i in (3, 4, 5):
    for j in range(5):
        set_cell(T3.rows[i].cells[j], "")

# ================= 附表：经费预算表 =================
T4 = tables[3]
rows4 = [
    ("600元", "专业资料购置300元；调研与结题材料打印装订300元"),
    ("0元", "利用现有开发设备，不新增仪器设备购置"),
    ("0元", "无相关费用支出"),
    ("900元", "模型调用与计算服务租赁暂按150元/月×6个月估算，按实际服务及校内口径执行"),
    ("0元", "无相关费用支出"),
    ("0元", "无相关费用支出"),
    ("500元", "软件著作权申请及知识产权事务费用预留500元，按实际需要支出"),
]
for k, (amt, why) in enumerate(rows4, start=1):
    set_cell(T4.rows[k].cells[2], amt)
    set_cell(T4.rows[k].cells[3], why)
set_cell(T4.rows[8].cells[1], "2000元")

# ================= 收尾：删多余空行、空表格行 =================
def del_row(t, idx):
    tr = t.rows[idx]._element
    tr.getparent().remove(tr)


# 1) 基本情况表里没用到的成员空行
for i in (12, 11, 10):
    del_row(tables[0], i)
# 2) 成员分工表里没用到的空行
for i in (5, 4, 3):
    del_row(tables[2], i)

# 3) 删除“三、经费概算”到“附表：”之间多余的空段落（保留带分节符的那段）
_ps = d.paragraphs
_i_a = next(i for i, p in enumerate(_ps) if p.text.strip().startswith("三、经费概算"))
_i_b = next(i for i, p in enumerate(_ps) if p.text.strip() == "附表：")
for _p in _ps[_i_a + 1:_i_b]:
    if _p.text.strip():
        continue
    _pPr = _p._element.find(qn("w:pPr"))
    if _pPr is not None and _pPr.find(qn("w:sectPr")) is not None:
        continue
    _p._element.getparent().remove(_p._element)

# 4) 文末多余空段落
for _p in list(d.paragraphs)[-3:]:
    if not _p.text.strip():
        _pPr = _p._element.find(qn("w:pPr"))
        if _pPr is None or _pPr.find(qn("w:sectPr")) is None:
            _p._element.getparent().remove(_p._element)

d.save(OUT)
print("SAVED:", OUT)
