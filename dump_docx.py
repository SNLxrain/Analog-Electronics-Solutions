# -*- coding: utf-8 -*-
"""导出 docx 的结构信息，用于仿照模板"""
import sys
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

path = sys.argv[1]
d = Document(path)

print("=" * 70)
print("SECTIONS / PAGE SETUP")
for i, s in enumerate(d.sections):
    print(f"[{i}] page {s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm "
          f"margins L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f} "
          f"T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f}")

print("=" * 70)
print("BODY (paragraphs + tables in document order)")
body = d.element.body
pi = ti = 0
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = d.paragraphs[pi]; pi += 1
        pf = p.paragraph_format
        fonts = []
        for r in p.runs[:3]:
            sz = r.font.size.pt if r.font.size else None
            ea = None
            if r._element.rPr is not None and r._element.rPr.rFonts is not None:
                ea = r._element.rPr.rFonts.get(qn('w:eastAsia'))
            fonts.append(f"(b={r.bold},sz={sz},font={r.font.name},ea={ea})")
        print(f"P{pi:03d} [{p.style.name}] align={pf.alignment} "
              f"first_line={pf.first_line_indent} line={pf.line_spacing} "
              f"{' '.join(fonts)}")
        if p.text.strip():
            print(f"      TEXT: {p.text}")
    elif tag == 'tbl':
        t = d.tables[ti]; ti += 1
        print(f"--- TABLE {ti}: {len(t.rows)} rows x {len(t.columns)} cols, style={t.style.name if t.style else None}")
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace('\n', '\\n') for c in row.cells]
            print(f"    R{ri}: {' | '.join(cells)}")
print("=" * 70)
print("ALL STYLE NAMES:", sorted({p.style.name for p in d.paragraphs}))
